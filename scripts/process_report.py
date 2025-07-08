from datetime import datetime

import os
from docx import Document

from scripts.extract_data_pdf import extract_data_from_pdf  
from scripts.generate_table_excel import insert_table_from_excel

def fill_word_template(template_path, output_path, data, excel_path):
    doc = Document(template_path)

    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(key, value)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for key, value in data.items():
                        if key in paragraph.text:
                            for run in paragraph.runs:
                                run.text = run.text.replace(key, value)

    insert_table_from_excel(doc, excel_path)

    doc.save(output_path)

def generate_report_from_data(pdf_path, excel_path=None, session_data=None):
    data = extract_data_from_pdf(pdf_path)
    if not data:
        return None
    
    data.update(session_data)

    template_path = os.path.join("doc_template", "template.docx")
    doc_name = f"Informe - {datetime.now().strftime("%Y-%m-%d_%H-%M-%S")}"
    output_path = os.path.join("output", f"{doc_name}.docx")

    try:
        fill_word_template(template_path, output_path, data, excel_path=excel_path)
        return output_path
    except Exception as e:
        return None

if __name__ == "__main__":
    pdf_path = "uploads/orden_ingreso.pdf"
    excel_path = "uploads/excel_datos.xlsx"
    fill_word_template(pdf_path, excel_path)
